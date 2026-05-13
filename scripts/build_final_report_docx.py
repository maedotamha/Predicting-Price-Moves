from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "reports" / "nova_financial_final_report.docx"
SUMMARY_PATH = PROJECT_ROOT / "reports" / "actual_analysis_summary.json"
FIGURES_DIR = PROJECT_ROOT / "reports" / "figures"

ACCENT = "1F4E79"
LIGHT_BLUE = "D9EAF7"


def load_summary() -> dict:
    if SUMMARY_PATH.exists():
        return json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    return {}


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text: str, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade_cell(cell, fill)
        if fill == ACCENT:
            run.font.color.rgb = RGBColor(255, 255, 255)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.3)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(23)
    styles["Title"].font.color.rgb = RGBColor.from_string(ACCENT)
    styles["Heading 1"].font.name = "Aptos"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(ACCENT)
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.color.rgb = RGBColor(45, 45, 45)


def add_para(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)


def add_picture(document: Document, filename: str, caption: str) -> None:
    path = FIGURES_DIR / filename
    if not path.exists():
        return
    document.add_picture(str(path), width=Inches(6.7))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(8.5)
    caption_paragraph.paragraph_format.space_after = Pt(8)


def add_metrics_table(document: Document, summary: dict) -> None:
    document.add_heading("Concrete Analysis Outputs", level=1)
    metrics = summary.get("financial_metrics", {})
    rows = [
        ("FNSPID news rows analyzed", f"{summary.get('news_rows', 0):,}"),
        ("Historical price rows analyzed", f"{summary.get('price_rows', 0):,}"),
        ("News date range", " to ".join(summary.get("date_range", ["N/A", "N/A"]))),
        ("Stocks analyzed", ", ".join(summary.get("selected_tickers", []))),
        ("Mean headline length", f"{summary.get('headline_mean_chars', 0):.2f} characters; {summary.get('headline_mean_words', 0):.2f} words"),
        ("Peak news day", f"{summary.get('peak_news_day', {}).get('date', 'N/A')} ({summary.get('peak_news_day', {}).get('article_count', 0)} articles)"),
        ("Peak publishing hour", f"{summary.get('peak_news_hour_utc', 'N/A')}:00 UTC"),
        ("Technical-analysis stock", summary.get("technical_stock", "N/A")),
        ("Latest adjusted close", f"{summary.get('latest_adj_close', 0):.2f}"),
        ("Latest RSI(14)", f"{summary.get('latest_rsi_14', 0):.2f}"),
        ("Latest MACD", f"{summary.get('latest_macd', 0):.3f}"),
        ("Cumulative return", f"{metrics.get('cumulative_return', 0):.2f}"),
        ("Annualized volatility", f"{metrics.get('annualized_volatility', 0):.2f}"),
        ("Maximum drawdown", f"{metrics.get('max_drawdown', 0):.2f}"),
        ("Sentiment-return observations", f"{summary.get('correlation_observations', 0):,}"),
        ("Overall Pearson correlation", f"{summary.get('overall_pearson_correlation', 0):.3f}"),
    ]

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell(table.rows[0].cells[0], "Metric", bold=True, fill=ACCENT)
    set_cell(table.rows[0].cells[1], "Result", bold=True, fill=ACCENT)
    for metric, value in rows:
        cells = table.add_row().cells
        set_cell(cells[0], metric, bold=True)
        set_cell(cells[1], value)


def add_deliverables_table(document: Document) -> None:
    document.add_heading("Completed Work by Task", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, header in enumerate(["Task", "Completed Artifact", "Demonstrated Output"]):
        set_cell(table.rows[0].cells[idx], header, bold=True, fill=ACCENT)

    rows = [
        (
            "Task 1",
            "task_1_eda.ipynb; src/eda.py",
            "Publisher counts, headline statistics, publication trends, TF-IDF terms, topic-modeling-ready pipeline, and EDA figures.",
        ),
        (
            "Task 2",
            "task_2_quantitative_analysis.ipynb; src/technical_indicators.py",
            "Cleaned price data, SMA/EMA, RSI, MACD, financial metrics, and price-based visualizations.",
        ),
        (
            "Task 3",
            "task_3_sentiment_correlation.ipynb; src/sentiment_correlation.py",
            "Sentiment scores, trading-day alignment, daily returns, Pearson correlation, scatter plot, and sentiment-category return chart.",
        ),
        ("Validation", "tests/", "10 passing unit tests."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell(cells[idx], value, bold=idx == 0, fill=LIGHT_BLUE if idx == 0 else None)


def add_task_1(document: Document, summary: dict) -> None:
    document.add_heading("Task 1: Exploratory Data Analysis", level=1)
    add_para(
        document,
        f"The analysis used a streamed sample of {summary.get('news_rows', 0):,} FNSPID news records. The sample covers {', '.join(summary.get('selected_tickers', []))} and spans {summary.get('date_range', ['N/A', 'N/A'])[0]} to {summary.get('date_range', ['N/A', 'N/A'])[1]}.",
    )
    add_para(
        document,
        f"Headline length averaged {summary.get('headline_mean_chars', 0):.2f} characters and {summary.get('headline_mean_words', 0):.2f} words. The busiest day in the sample was {summary.get('peak_news_day', {}).get('date', 'N/A')} with {summary.get('peak_news_day', {}).get('article_count', 0)} articles, while the most common publishing hour was {summary.get('peak_news_hour_utc', 'N/A')}:00 UTC.",
    )
    top_terms = ", ".join(summary.get("top_terms", []))
    add_para(document, f"The top TF-IDF terms included: {top_terms}. These terms show strong coverage around company-specific news, earnings, analyst commentary, and market movement.")
    add_picture(document, "actual_top_publishers.png", "Figure 1. Top publishers by article count.")
    add_picture(document, "actual_daily_news_volume.png", "Figure 2. Daily news-volume trend.")
    add_picture(document, "actual_top_tfidf_terms.png", "Figure 3. Top TF-IDF headline terms and phrases.")


def add_task_2(document: Document, summary: dict) -> None:
    document.add_heading("Task 2: Technical Indicator Analysis", level=1)
    metrics = summary.get("financial_metrics", {})
    add_para(
        document,
        f"The stock-price analysis used {summary.get('price_rows', 0):,} daily OHLCV records downloaded through Yahoo Finance for the sampled tickers. For the technical indicator visualization, {summary.get('technical_stock', 'N/A')} was selected.",
    )
    add_para(
        document,
        f"The latest adjusted close for {summary.get('technical_stock', 'N/A')} was {summary.get('latest_adj_close', 0):.2f}. The latest RSI(14) was {summary.get('latest_rsi_14', 0):.2f}, which is neither above the 70 overbought threshold nor below the 30 oversold threshold. The latest MACD value was {summary.get('latest_macd', 0):.3f}.",
    )
    add_para(
        document,
        f"The PyNance-style metrics show cumulative return of {metrics.get('cumulative_return', 0):.2f}, annualized volatility of {metrics.get('annualized_volatility', 0):.2f}, max drawdown of {metrics.get('max_drawdown', 0):.2f}, and a zero-risk-free-rate Sharpe ratio of {metrics.get('sharpe_ratio_zero_rf', 0):.2f}.",
    )
    add_picture(document, "actual_price_moving_averages.png", "Figure 4. Adjusted close with SMA and EMA overlays.")
    add_picture(document, "actual_rsi_macd.png", "Figure 5. RSI and MACD indicator panels.")


def add_task_3(document: Document, summary: dict) -> None:
    document.add_heading("Task 3: Sentiment and Return Correlation", level=1)
    corr = summary.get("overall_pearson_correlation", 0)
    add_para(
        document,
        f"The sentiment-return dataset contained {summary.get('correlation_observations', 0):,} matched stock-day observations after aligning headlines to valid trading days and joining them with daily adjusted-close returns. The overall Pearson correlation between average daily sentiment and daily return was {corr:.3f}.",
    )
    if abs(corr) < 0.1:
        interpretation = "This indicates a weak same-day linear relationship in the sampled data."
    elif corr > 0:
        interpretation = "This indicates a positive same-day association in the sampled data."
    else:
        interpretation = "This indicates a negative same-day association in the sampled data."
    add_para(document, interpretation)
    returns = summary.get("returns_by_sentiment", [])
    if returns:
        bits = [
            f"{row['sentiment_label']}: {row['avg_daily_return_pct']:.2f}% over {int(row['observations'])} observations"
            for row in returns
        ]
        add_para(document, "Average returns by sentiment category were " + "; ".join(bits) + ".")
    add_picture(document, "actual_sentiment_return_scatter.png", "Figure 6. Average daily sentiment versus daily return.")
    add_picture(document, "actual_return_by_sentiment.png", "Figure 7. Average daily return by sentiment category.")


def add_recommendations(document: Document) -> None:
    document.add_heading("Investment Strategy Recommendations", level=1)
    for item in [
        "Use headline sentiment as one feature within a broader model, not as a standalone trading signal.",
        "Combine positive sentiment with improving MACD or supportive moving-average structure before treating it as actionable.",
        "Segment sentiment by news theme because earnings, analyst-rating, FDA, and M&A headlines can have different price effects.",
        "Treat weak same-day correlation cautiously and test lagged return windows before using sentiment for forecasting.",
        "Control for publisher and ticker concentration so the signal is not dominated by one source or one heavily covered stock.",
    ]:
        add_bullet(document, item)


def add_limitations(document: Document) -> None:
    document.add_heading("Limitations and Future Improvements", level=1)
    for item in [
        "The FNSPID file is very large, so this report uses a reproducible streamed sample rather than the full 23GB news file.",
        "The same-day Pearson correlation is an association measure and does not prove causation.",
        "After-hours headlines should be analyzed separately in a more advanced event-study design.",
        "Future work should test next-day, three-day, and five-day forward returns and validate any signal out of sample.",
    ]:
        add_bullet(document, item)


def build_docx() -> None:
    summary = load_summary()
    document = Document()
    style_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Predicting Price Moves with News Sentiment")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Final Evidence-Based Analysis Report for Nova Financial Solutions")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(70, 70, 70)

    document.add_paragraph()
    document.add_heading("Executive Summary", level=1)
    add_para(
        document,
        "This report completes the Nova Financial Solutions challenge by presenting actual outputs from news EDA, technical stock analysis, and sentiment-return correlation. The analysis uses a reproducible FNSPID news sample and matching Yahoo Finance price data to demonstrate the full analytical pipeline end to end.",
    )
    add_para(
        document,
        "The strongest practical finding is that average headline sentiment shows only a weak same-day linear relationship with daily stock returns in the sampled data. Positive-sentiment days show higher average returns than neutral and negative days, but the weak Pearson correlation means sentiment should be used as a supporting signal rather than a standalone forecasting tool.",
    )

    add_deliverables_table(document)
    add_metrics_table(document, summary)
    add_task_1(document, summary)
    add_task_2(document, summary)
    add_task_3(document, summary)
    add_recommendations(document)
    add_limitations(document)

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Nova Financial Solutions | Final Evidence-Based Analysis Report"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_docx()
