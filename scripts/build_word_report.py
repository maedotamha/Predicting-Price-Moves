from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "reports" / "nova_financial_task_1_2_report.docx"


ACCENT = "1F4E79"
LIGHT_ACCENT = "D9EAF7"
SOFT_GRAY = "F2F5F7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_document(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    for style_name, size, color in [
        ("Title", 24, ACCENT),
        ("Heading 1", 16, ACCENT),
        ("Heading 2", 13, "2B2B2B"),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    for section in document.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)


def add_section(document: Document, title: str, body: list[str]) -> None:
    document.add_heading(title, level=1)
    for item in body:
        paragraph = document.add_paragraph(item)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.08


def add_status_table(document: Document) -> None:
    document.add_heading("Repository Status", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Area", "Status", "Notes"]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], ACCENT)
        table.rows[0].cells[index].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    rows = [
        ("Task 1", "Complete", "EDA notebook, helper modules, CI, and tests are in place."),
        ("Task 2", "Complete", "Price-data loader, indicators, metrics, notebook, and tests are in place."),
        ("Raw data", "Pending", "Run notebooks after adding news and stock-price CSV files to data/raw/."),
        ("Validation", "Passed", "Unit test suite passed locally with 7 tests."),
    ]
    widths = [1.25, 1.1, 4.8]
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value, bold=index == 1)
            if index == 1:
                set_cell_shading(cells[index], LIGHT_ACCENT if value != "Pending" else SOFT_GRAY)
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def add_recommendations(document: Document) -> None:
    document.add_heading("Strategic Recommendations", level=1)
    recommendations = [
        ("Use adjusted close returns", "Adjusted prices account for splits and dividends, making them more reliable for historical return analysis."),
        ("Align news to trading sessions", "After-hours headlines should generally be mapped to the next trading session when measuring price impact."),
        ("Segment sentiment by topic", "Earnings, analyst ratings, FDA events, and M&A headlines may have different return relationships."),
        ("Control for publisher concentration", "Evaluate whether sentiment patterns reflect market narratives or repeated editorial style."),
        ("Combine sentiment with indicators", "Signals may be stronger when positive sentiment aligns with improving momentum indicators."),
        ("Measure multiple event windows", "Compare same-day, next-day, and multi-day returns after publication."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "Recommendation", bold=True)
    set_cell_text(table.rows[0].cells[1], "Rationale", bold=True)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, ACCENT)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    for title, rationale in recommendations:
        cells = table.add_row().cells
        set_cell_text(cells[0], title, bold=True)
        set_cell_text(cells[1], rationale)
    for row in table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.9)


def build_report() -> None:
    document = Document()
    style_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Nova Financial Solutions")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Financial News and Price-Move Analysis Report")
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(60, 60, 60)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("Tasks 1 and 2 | Prepared for the financial news sentiment challenge")
    meta_run.font.size = Pt(10.5)

    document.add_paragraph()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "This report summarizes the first two stages of a financial intelligence pipeline for Nova Financial Solutions. "
        "Task 1 establishes exploratory analysis for financial-news headlines, publishers, publication timing, and recurring topics. "
        "Task 2 extends the workflow into quantitative market analysis by loading historical OHLCV price data, cleaning it, "
        "computing technical indicators, and visualizing price behavior."
    )
    document.add_paragraph(
        "The repository now contains reproducible notebooks, tested helper modules, CI configuration, and clean Git branches. "
        "The analysis is ready to run once the raw FNSPID news dataset and historical stock-price CSV files are placed in data/raw/."
    )

    add_status_table(document)

    add_section(
        document,
        "Business Objective",
        [
            "Nova Financial Solutions wants to improve forecasting accuracy by connecting market narratives to price action. "
            "The core analytical question is whether financial-news sentiment and publishing patterns can help explain or anticipate stock-price movement.",
            "The current project builds the foundation for that analysis by preparing news EDA, keyword discovery, publisher analysis, stock-price cleaning, technical indicators, and risk-return metrics.",
        ],
    )

    document.add_heading("Task 1: Exploratory Data Analysis", level=1)
    for item in [
        "The news loader expects headline, url, publisher, date, and stock fields.",
        "Publication timestamps are parsed as UTC-aware dates, and ticker symbols are normalized.",
        "The notebook computes headline character and word counts, article counts by publisher, article counts by ticker, and publication date ranges.",
        "TF-IDF is used to identify high-signal headline terms and phrases.",
        "LDA topic modeling provides an initial view of recurring themes such as earnings, ratings, FDA events, M&A, or guidance updates.",
        "Publication volume is analyzed by calendar date and UTC hour to identify spikes and timing patterns.",
        "Publisher email domains are extracted when available to understand organizational contribution patterns.",
    ]:
        add_bullet(document, item)

    document.add_heading("Task 2: Quantitative Price Analysis", level=1)
    for item in [
        "The price loader expects Date, Open, High, Low, Close, and Volume; Adj Close and Stock are recommended.",
        "OHLCV fields are converted to numeric values, dates are parsed, and adjusted close is used for return calculations when available.",
        "Missing numeric values are forward/backward filled within each ticker, then unusable rows are removed.",
        "The pipeline computes SMA, EMA, RSI, MACD, MACD signal, MACD histogram, and daily adjusted-close return.",
        "TA-Lib is used when installed, with pandas fallbacks to keep the notebook executable in environments without TA-Lib binaries.",
        "PyNance availability is recorded, while stable metrics are computed for cumulative return, annualized volatility, maximum drawdown, and Sharpe ratio.",
        "The notebook saves price-with-moving-average, RSI, and MACD visualizations to reports/figures/ when executed.",
    ]:
        add_bullet(document, item)

    add_section(
        document,
        "Data Quality Considerations",
        [
            "The raw FNSPID and historical price datasets are not included in the repository, so this report does not claim dataset-specific numeric findings.",
            "Before interpreting results, confirm ticker alignment, remove duplicate ticker-date rows, verify adjusted-close quality, and map news timestamps to the correct trading sessions.",
            "Publisher concentration and topic mix should be checked before treating aggregate sentiment as a single predictive signal.",
        ],
    )

    add_recommendations(document)

    add_section(
        document,
        "Limitations",
        [
            "The current work creates the analytical foundation but does not yet produce final investment conclusions because raw data is not present in the workspace.",
            "Technical indicators are descriptive tools, not standalone forecasts. Sentiment-return correlations should be tested out of sample before being used in an investment strategy.",
        ],
    )

    document.add_heading("Next Steps", level=1)
    for item in [
        "Add the FNSPID news CSV to data/raw/.",
        "Add historical stock-price CSV data to data/raw/.",
        "Run notebooks/task_1_eda.ipynb.",
        "Run notebooks/task_2_quantitative_analysis.ipynb.",
        "Review generated figures in reports/figures/.",
        "Proceed to sentiment scoring and correlation analysis in Task 3.",
    ]:
        add_bullet(document, item)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Nova Financial Solutions | Financial News and Price-Move Analysis"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()
